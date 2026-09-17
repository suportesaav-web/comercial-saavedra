import os
import openpyxl
import pandas as pd
import numpy as np
import json
import re
import datetime as dt

file_path = os.path.join('dados', 'bruto', 'Tarefas Power BI.xlsx')
wb = openpyxl.load_workbook(file_path, data_only=True)
ws = wb['Ploomes']

# Raw types from openpyxl
raw_header = [cell.value for cell in ws[1]]

raw_types = {}
for col_idx, col_name in enumerate(raw_header, start=1):
    types_in_col = set()
    for row in range(2, ws.max_row + 1):
        v = ws.cell(row=row, column=col_idx).value
        if v is not None:
            types_in_col.add(type(v).__name__)
        else:
            types_in_col.add('NoneType')
    raw_types[col_name] = list(types_in_col)

# Read into pandas
df = pd.read_excel(file_path, sheet_name='Ploomes')

report = {}
report['file_info'] = {
    'file_name': os.path.basename(file_path),
    'extension': os.path.splitext(file_path)[1],
    'size_bytes': os.path.getsize(file_path),
    'sheet_count': len(wb.sheetnames),
    'sheets': wb.sheetnames,
    'total_rows_excel': ws.max_row,
    'total_cols_excel': ws.max_column,
    'data_rows': len(df),
    'data_cols': len(df.columns),
}

# Column profiling
col_profiles = {}
for col in df.columns:
    s = df[col]
    val_count = int(s.count())
    null_count = int(s.isnull().sum())
    null_pct = float(null_count / len(df) * 100)
    distinct_count = int(s.nunique(dropna=True))
    cardinality_pct = float(distinct_count / val_count * 100) if val_count > 0 else 0.0
    
    col_info = {
        'column_name': col,
        'inferred_type_pandas': str(s.dtype),
        'real_storage_types_excel': raw_types.get(col, []),
        'total_rows': len(df),
        'valid_values_count': val_count,
        'null_count': null_count,
        'null_pct': round(null_pct, 2),
        'distinct_count': distinct_count,
        'cardinality_pct': round(cardinality_pct, 2),
    }
    
    # Specific type analysis
    if pd.api.types.is_datetime64_any_dtype(s):
        dt_s = s.dropna()
        col_info['type_family'] = 'datetime'
        col_info['min'] = str(dt_s.min()) if not dt_s.empty else None
        col_info['max'] = str(dt_s.max()) if not dt_s.empty else None
        
        # Check time presence
        has_time = dt_s.apply(lambda x: x.hour != 0 or x.minute != 0 or x.second != 0).any()
        col_info['has_time'] = bool(has_time)
        
        # Gaps / distribution
        col_info['years'] = sorted([int(y) for y in dt_s.dt.year.unique()])
        col_info['future_dates_count'] = int((dt_s > pd.Timestamp.now()).sum())
        
    elif pd.api.types.is_bool_dtype(s) or (set(s.dropna().unique()).issubset({True, False})):
        col_info['type_family'] = 'boolean'
        col_info['value_counts'] = {str(k): int(v) for k, v in s.value_counts(dropna=False).items()}
        col_info['min'] = str(s.min())
        col_info['max'] = str(s.max())
        
    elif pd.api.types.is_numeric_dtype(s):
        num_s = s.dropna()
        col_info['type_family'] = 'numeric'
        col_info['min'] = float(num_s.min()) if not num_s.empty else None
        col_info['max'] = float(num_s.max()) if not num_s.empty else None
        col_info['mean'] = float(num_s.mean()) if not num_s.empty else None
        col_info['median'] = float(num_s.median()) if not num_s.empty else None
        col_info['zero_count'] = int((num_s == 0).sum())
        col_info['negative_count'] = int((num_s < 0).sum())
        
    else: # string / object
        str_s = s.dropna().astype(str)
        col_info['type_family'] = 'string'
        col_info['min_len'] = int(str_s.str.len().min()) if not str_s.empty else 0
        col_info['max_len'] = int(str_s.str.len().max()) if not str_s.empty else 0
        
        # Text quality checks
        leading_trailing_ws = int(str_s.apply(lambda x: len(x) != len(x.strip())).sum())
        col_info['leading_trailing_whitespace_count'] = leading_trailing_ws
        
        # Check case inconsistencies
        lower_distinct = str_s.str.lower().nunique()
        col_info['case_inconsistency_potential'] = bool(lower_distinct < distinct_count)
        col_info['lower_distinct_count'] = int(lower_distinct)
        
        # Sample values
        top_vals = str_s.value_counts().head(5).to_dict()
        col_info['top_5_values'] = {k: int(v) for k, v in top_vals.items()}
        
    # Sample non-null values
    col_info['sample_values'] = s.dropna().head(5).astype(str).tolist()
    col_profiles[col] = col_info

report['columns'] = col_profiles

# Key candidates testing
key_tests = []
for col in df.columns:
    val_c = len(df)
    dist_c = int(df[col].nunique(dropna=False))
    null_c = int(df[col].isnull().sum())
    is_unique = (dist_c == val_c) and (null_c == 0)
    key_tests.append({
        'columns': [col],
        'total_rows': val_c,
        'distinct_count': dist_c,
        'null_count': null_c,
        'uniqueness_pct': round(dist_c / val_c * 100, 2),
        'is_primary_key_candidate': is_unique
    })

combo_candidates = [
    ['Data', 'Usuários'],
    ['Data', 'Criador'],
    ['Data', 'Nome do Cliente'],
    ['Título', 'Data'],
    ['Título', 'Nome do Cliente', 'Data'],
    ['Título', 'Usuários', 'Data'],
    ['Data de criação', 'Usuários'],
    ['Data de criação', 'Criador'],
    ['Data de criação', 'Título']
]

for combo in combo_candidates:
    subset = df[combo]
    val_c = len(df)
    dist_c = len(subset.drop_duplicates())
    null_c = int(subset.isnull().any(axis=1).sum())
    key_tests.append({
        'columns': combo,
        'total_rows': val_c,
        'distinct_count': dist_c,
        'null_count': null_c,
        'uniqueness_pct': round(dist_c / val_c * 100, 2),
        'is_primary_key_candidate': (dist_c == val_c) and (null_c == 0)
    })

# Check full duplicates
full_dups = int(df.duplicated(keep=False).sum())
report['full_duplicates_count'] = full_dups
if full_dups > 0:
    report['full_duplicates_samples'] = df[df.duplicated(keep=False)].head(6).to_dict(orient='records')
else:
    report['full_duplicates_samples'] = []

cols_no_ts = [c for c in df.columns if c not in ['Data de criação']]
report['duplicates_excluding_data_criacao'] = int(df.duplicated(subset=cols_no_ts, keep=False).sum())

report['key_tests'] = key_tests

with open(os.path.join('documentacao', 'scripts_analise', 'resultado_perfil.json'), 'w', encoding='utf-8') as f:
    json.dump(report, f, indent=2, default=str, ensure_ascii=False)

print("Análise preliminar salva com sucesso em documentacao/scripts_analise/resultado_perfil.json")
