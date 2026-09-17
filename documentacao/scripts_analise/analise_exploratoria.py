import os
import openpyxl
import pandas as pd
import numpy as np
import json
from datetime import datetime

file_path = os.path.join('dados', 'bruto', 'Tarefas Power BI.xlsx')

print(f"=== ANÁLISE DO ARQUIVO: {file_path} ===")
file_size = os.path.getsize(file_path)
print(f"Tamanho do arquivo: {file_size} bytes ({file_size / 1024:.2f} KB)")

# Open with openpyxl to check formulas, merged cells, hidden rows/cols
wb_formula = openpyxl.load_workbook(file_path, data_only=False)
wb_data = openpyxl.load_workbook(file_path, data_only=True)

sheet_names = wb_formula.sheetnames
print(f"Número de abas: {len(sheet_names)}")
print(f"Nomes das abas: {sheet_names}")

inventory = {}

for sname in sheet_names:
    ws_f = wb_formula[sname]
    ws_d = wb_data[sname]
    
    merged_ranges = [str(r) for r in ws_f.merged_cells.ranges]
    
    hidden_rows = [r for r, dim in ws_f.row_dimensions.items() if dim.hidden]
    hidden_cols = [c for c, dim in ws_f.column_dimensions.items() if dim.hidden]
    
    # Formulas check
    formula_cells = []
    for row in ws_f.iter_rows(values_only=False):
        for cell in row:
            if cell.data_type == 'f' or (isinstance(cell.value, str) and str(cell.value).startswith('=')):
                formula_cells.append((cell.coordinate, cell.value))
                
    inventory[sname] = {
        'max_row': ws_f.max_row,
        'max_column': ws_f.max_column,
        'merged_cells_count': len(merged_ranges),
        'merged_cells': merged_ranges,
        'hidden_rows_count': len(hidden_rows),
        'hidden_rows': hidden_rows,
        'hidden_cols_count': len(hidden_cols),
        'hidden_cols': hidden_cols,
        'formula_cells_count': len(formula_cells),
        'formula_samples': formula_cells[:10]
    }
    
    print(f"\nAba: {sname}")
    print(f"  Linhas max: {ws_f.max_row}, Colunas max: {ws_f.max_column}")
    print(f"  Células mescladas: {len(merged_ranges)}")
    print(f"  Linhas ocultas: {len(hidden_rows)}")
    print(f"  Colunas ocultas: {len(hidden_cols)}")
    print(f"  Células com fórmulas: {len(formula_cells)}")

# Read with Pandas
df = pd.read_excel(file_path, sheet_name=sheet_names[0])
print(f"\nDataFrame lido da aba '{sheet_names[0]}':")
print(f"Shape: {df.shape} (Linhas: {df.shape[0]}, Colunas: {df.shape[1]})")
print("\nColunas encontradas:")
for idx, col in enumerate(df.columns):
    print(f"  [{idx}] '{col}' - Tipo Pandas: {df[col].dtype}")

print("\nPrimeiras 3 linhas:")
print(df.head(3).to_dict(orient='records'))

print("\nVerificando linhas vazias:")
empty_rows = df.isnull().all(axis=1).sum()
print(f"Linhas totalmente vazias no DataFrame: {empty_rows}")

empty_cols = df.isnull().all(axis=0).sum()
print(f"Colunas totalmente vazias no DataFrame: {empty_cols}")

